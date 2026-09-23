#!/usr/bin/env python3
"""Formüllü DCF (FCFF) Excel modeli üretir.

Kullanım:  python3 scripts/dcf_excel.py girdi.json cikti.xlsx

Girdi (JSON; oranlar ondalık: 0.25 = %25):
{
  "sirket": "THYAO", "para_birimi": "TL", "birim": "milyon",
  "baz_yil": 2025, "baz_satis": 100000,
  "buyume": [0.30, 0.25, 0.20, 0.15, 0.12],   # yıl yıl satış büyümesi (liste uzunluğu = projeksiyon yılı)
  "ebit_marji": [0.12, 0.12, 0.13, 0.13, 0.13],  # tek sayı da olabilir
  "vergi": 0.25,
  "satis_sermaye": 2.0,          # yeniden yatırım = satış artışı / satış-sermaye oranı
  "wacc": 0.30,                  # tek sayı veya yıl yıl liste
  "terminal_buyume": 0.05, "terminal_wacc": 0.20,   # terminal_wacc boşsa son yılın WACC'ı
  "net_borc": 20000, "azinlik": 0, "istirakler": 0,
  "hisse_sayisi": 1380, "guncel_fiyat": 300
}
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(bold=True, color="FFFFFF")
INPUT = PatternFill("solid", fgColor="FFF4D6")
PCT = "0.0%"
NUM = "#,##0"
NUM2 = "#,##0.00"


def as_list(v, n):
    return list(v) if isinstance(v, list) else [v] * n


def model_values(a):
    """Duyarlılık tablosu için aynı modelin Python hesabı (Excel formülleriyle birebir)."""
    n = len(a["buyume"])
    g, m, w = a["buyume"], as_list(a["ebit_marji"], n), as_list(a["wacc"], n)

    def value(wacc_shift=0.0, tg=None):
        tg = a["terminal_buyume"] if tg is None else tg
        sales, prev, disc, pv = a["baz_satis"], a["baz_satis"], 1.0, 0.0
        fcff = 0.0
        for i in range(n):
            sales = prev * (1 + g[i])
            nopat = sales * m[i] * (1 - a["vergi"])
            fcff = nopat - (sales - prev) / a["satis_sermaye"]
            disc /= 1 + w[i] + wacc_shift
            pv += fcff * disc
            prev = sales
        tw = (a.get("terminal_wacc") or w[-1]) + wacc_shift
        tv = fcff * (1 + tg) / (tw - tg)
        equity = pv + tv * disc - a["net_borc"] - a.get("azinlik", 0) + a.get("istirakler", 0)
        return equity / a["hisse_sayisi"]

    return value


def build(a, out):
    n = len(a["buyume"])
    wb = Workbook()

    # ── Varsayımlar ──
    ws = wb.active
    ws.title = "Varsayımlar"
    ws["A1"] = f"{a.get('sirket', '')} DCF varsayımları ({a.get('birim', 'milyon')} {a.get('para_birimi', 'TL')})"
    ws["A1"].font = Font(bold=True, size=13)
    rows = [
        ("Baz yıl", a["baz_yil"], None), ("Baz yıl satış", a["baz_satis"], NUM), ("Vergi oranı", a["vergi"], PCT),
        ("Satış / sermaye", a["satis_sermaye"], NUM2), ("Terminal büyüme", a["terminal_buyume"], PCT),
        ("Terminal WACC", a.get("terminal_wacc") or as_list(a["wacc"], n)[-1], PCT), ("Net borç", a["net_borc"], NUM),
        ("Azınlık payları", a.get("azinlik", 0), NUM), ("İştirakler (+)", a.get("istirakler", 0), NUM),
        ("Hisse sayısı", a["hisse_sayisi"], NUM2), ("Güncel fiyat", a.get("guncel_fiyat", 0), NUM2),
    ]
    ref = {}
    for i, (label, val, fmt) in enumerate(rows, start=3):
        ws.cell(i, 1, label)
        c = ws.cell(i, 2, val)
        c.fill = INPUT
        if fmt:
            c.number_format = fmt
        ref[label] = f"Varsayımlar!$B${i}"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    r0 = len(rows) + 5
    ws.cell(r0, 1, "Yıl yıl").font = BOLD
    ws.cell(r0 + 1, 1, "Satış büyümesi")
    ws.cell(r0 + 2, 1, "EBIT marjı")
    ws.cell(r0 + 3, 1, "WACC")
    for j, (gv, mv, wv) in enumerate(zip(a["buyume"], as_list(a["ebit_marji"], n), as_list(a["wacc"], n)), start=2):
        ws.cell(r0, j, a["baz_yil"] + j - 1).font = BOLD
        for k, v in enumerate((gv, mv, wv), start=1):
            c = ws.cell(r0 + k, j, v)
            c.fill = INPUT
            c.number_format = PCT
    ws.cell(r0 + 5, 1, "Sarı hücreler girdidir; DCF sayfası formüllerle güncellenir.").font = Font(italic=True, color="6B7280")

    # ── DCF ──
    d = wb.create_sheet("DCF")
    labels = ["Satış", "Büyüme", "EBIT marjı", "EBIT", "Vergi", "NOPAT", "Satış artışı", "Yeniden yatırım", "FCFF", "WACC", "İskonto faktörü", "FCFF bugünkü değeri"]
    d.cell(1, 1, "Kalem").font = HEAD_FONT
    d.cell(1, 1).fill = HEAD
    for i, lab in enumerate(labels, start=2):
        d.cell(i, 1, lab)
    d.cell(1, 2, a["baz_yil"]).font = HEAD_FONT
    d.cell(1, 2).fill = HEAD
    d.cell(2, 2, f"={ref['Baz yıl satış']}").number_format = NUM
    for j in range(3, n + 3):
        col, prev = get_column_letter(j), get_column_letter(j - 1)
        src = get_column_letter(j - 1)  # Varsayımlar sayfasında aynı yılın sütunu
        d.cell(1, j, a["baz_yil"] + j - 2).font = HEAD_FONT
        d.cell(1, j).fill = HEAD
        f = {
            2: f"={prev}2*(1+{col}3)", 3: f"=Varsayımlar!{src}{r0 + 1}", 4: f"=Varsayımlar!{src}{r0 + 2}", 5: f"={col}2*{col}4",
            6: f"={col}5*{ref['Vergi oranı']}", 7: f"={col}5-{col}6", 8: f"={col}2-{prev}2", 9: f"={col}8/{ref['Satış / sermaye']}",
            10: f"={col}7-{col}9", 11: f"=Varsayımlar!{src}{r0 + 3}",
            12: (f"={prev}12/(1+{col}11)" if j > 3 else f"=1/(1+{col}11)"), 13: f"={col}10*{col}12",
        }
        for r, formula in f.items():
            c = d.cell(r, j, formula)
            c.number_format = PCT if r in (3, 4, 11) else ("0.0000" if r == 12 else NUM)
    last = get_column_letter(n + 2)
    first = "C"
    s = 16
    summary = [
        ("FCFF bugünkü değerleri toplamı", f"=SUM({first}13:{last}13)", NUM),
        ("Terminal değer (son yıl sonu)", f"={last}10*(1+{ref['Terminal büyüme']})/({ref['Terminal WACC']}-{ref['Terminal büyüme']})", NUM),
        ("Terminal değerin bugünkü değeri", f"=B{s + 1}*{last}12", NUM),
        ("Firma değeri", f"=B{s}+B{s + 2}", NUM),
        ("Net borç (−)", f"={ref['Net borç']}", NUM),
        ("Azınlık payları (−)", f"={ref['Azınlık payları']}", NUM),
        ("İştirakler (+)", f"={ref['İştirakler (+)']}", NUM),
        ("Özkaynak değeri", f"=B{s + 3}-B{s + 4}-B{s + 5}+B{s + 6}", NUM),
        ("Hisse başı değer", f"=B{s + 7}/{ref['Hisse sayısı']}", NUM2),
        ("Güncel fiyat", f"={ref['Güncel fiyat']}", NUM2),
        ("Potansiyel", f"=IF(B{s + 9}>0,B{s + 8}/B{s + 9}-1,\"\")", PCT),
        ("Terminal değerin firma değerindeki payı", f"=B{s + 2}/B{s + 3}", PCT),
    ]
    for i, (lab, formula, fmt) in enumerate(summary):
        d.cell(s + i, 1, lab).font = BOLD if lab in ("Firma değeri", "Özkaynak değeri", "Hisse başı değer") else Font()
        d.cell(s + i, 2, formula).number_format = fmt
    d.column_dimensions["A"].width = 36
    for j in range(2, n + 3):
        d.column_dimensions[get_column_letter(j)].width = 13

    # ── Duyarlılık ──
    sv = wb.create_sheet("Duyarlılık")
    value = model_values(a)
    sv["A1"] = "Hisse başı değer: WACC değişimi × terminal büyüme (değerler; varsayım değişirse betiği yeniden çalıştırın)"
    sv["A1"].font = BOLD
    shifts = [-0.04, -0.02, 0.0, 0.02, 0.04]
    tgs = [a["terminal_buyume"] + x for x in (-0.02, -0.01, 0.0, 0.01, 0.02)]
    sv.cell(3, 1, "WACC Δ \\ g").font = BOLD
    for j, tg in enumerate(tgs, start=2):
        sv.cell(3, j, tg).number_format = PCT
        sv.cell(3, j).font = BOLD
    for i, sh in enumerate(shifts, start=4):
        sv.cell(i, 1, sh).number_format = "+0.0%;-0.0%;0.0%"
        sv.cell(i, 1).font = BOLD
        for j, tg in enumerate(tgs, start=2):
            c = sv.cell(i, j, value(sh, tg))
            c.number_format = NUM2
            c.alignment = Alignment(horizontal="right")
    sv.column_dimensions["A"].width = 14

    wb.save(out)
    return value()


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Kullanım: python3 scripts/dcf_excel.py girdi.json cikti.xlsx")
    with open(sys.argv[1], encoding="utf-8") as fh:
        args = json.load(fh)
    per_share = build(args, sys.argv[2])
    print(json.dumps({"cikti": sys.argv[2], "hisse_basi_deger": round(per_share, 2)}, ensure_ascii=False))
